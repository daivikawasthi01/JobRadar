"""Tests for deduplicated workbook output."""
import openpyxl

from src.filters import FilterResult
from src.models import Company, Job
from src.store import Store
from src.export import export_excel


def test_export_shows_one_cross_source_copy(tmp_path):
    store = Store(":memory:")
    companies = [
        Company("Example Inc.", "greenhouse", "example", 1),
        Company("Example", "lever", "example", 1),
    ]
    store.sync_companies(companies)
    for company, job_id, location in (
        (companies[0], "gh-1", "Bengaluru, India"),
        (companies[1], "lever-1", "Bengaluru India"),
    ):
        job = Job(
            job_id=f"{company.ats}:{company.slug}:{job_id}",
            company_key=company.key,
            company=company.name,
            ats=company.ats,
            title="Software Engineer, Intern",
            location=location,
            url=f"https://example.com/{job_id}",
            description="",
            posted_at=None,
            salary=None,
            remote=None,
        )
        store.apply_crawl(
            company,
            [(job, FilterResult(True, country="India", min_exp=0, exp_signal="title_entry"))],
            "2026-10-15T00:00:00Z",
        )

    workbook_path = tmp_path / "jobs.xlsx"
    export_excel(store, str(workbook_path))
    workbook = openpyxl.load_workbook(workbook_path, read_only=True)
    assert workbook["Matches"].max_row == 2
    assert workbook["All Open"].max_row == 2
    workbook.close()
    store.close()
