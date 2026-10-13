"""Excel exporter using pandas and openpyxl with multiple sheets and formatting."""
from datetime import datetime, timedelta, timezone
import logging
import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
from src.notify import IST, utc_to_ist_str
from src.store import Store

logger = logging.getLogger(__name__)


def export_excel(store: Store, path: str = "data/jobs.xlsx") -> None:
    """Export jobs and crawl statistics to Excel workbook atomically.

    Sheets:
    1. Matches: matched=1 AND status='open', newest first_seen first
    2. All Open: status='open'
    3. Stats: Summary metrics
    """
    out_path = Path(path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = out_path.with_suffix(".tmp.xlsx")

    cur = store.conn.cursor()

    # 1. Matches data
    matches_query = """
        SELECT
            title AS "Role",
            company AS "Company",
            location AS "Location",
            country AS "Country",
            min_exp AS "Min Exp",
            max_exp AS "Max Exp",
            exp_signal AS "Exp Signal",
            posted_at AS "Posted",
            first_seen AS "First Seen (IST)",
            salary AS "Salary",
            ats AS "Source",
            url AS "Link"
        FROM jobs
        WHERE matched = 1 AND status = 'open'
        ORDER BY first_seen DESC
    """
    matches_rows = cur.execute(matches_query).fetchall()

    # 2. All Open data
    all_open_query = """
        SELECT
            title AS "Role",
            company AS "Company",
            location AS "Location",
            country AS "Country",
            min_exp AS "Min Exp",
            max_exp AS "Max Exp",
            exp_signal AS "Exp Signal",
            posted_at AS "Posted",
            first_seen AS "First Seen (IST)",
            salary AS "Salary",
            ats AS "Source",
            url AS "Link",
            matched AS "Matched",
            reject_reason AS "Reject Reason"
        FROM jobs
        WHERE status = 'open'
        ORDER BY first_seen DESC
    """
    all_open_rows = cur.execute(all_open_query).fetchall()

    def _rows_to_dicts(rows: list, is_matches: bool = True) -> list[dict]:
        res = []
        for r in rows:
            d = dict(r)
            d["First Seen (IST)"] = utc_to_ist_str(d["First Seen (IST)"])
            if d.get("Posted"):
                d["Posted"] = utc_to_ist_str(d["Posted"])
            res.append(d)
        return res

    matches_list = _rows_to_dicts(matches_rows, is_matches=True)
    all_open_list = _rows_to_dicts(all_open_rows, is_matches=False)

    df_matches = pd.DataFrame(matches_list) if matches_list else pd.DataFrame(columns=[
        "Role", "Company", "Location", "Country", "Min Exp", "Max Exp",
        "Exp Signal", "Posted", "First Seen (IST)", "Salary", "Source", "Link"
    ])

    df_all_open = pd.DataFrame(all_open_list) if all_open_list else pd.DataFrame(columns=[
        "Role", "Company", "Location", "Country", "Min Exp", "Max Exp",
        "Exp Signal", "Posted", "First Seen (IST)", "Salary", "Source", "Link",
        "Matched", "Reject Reason"
    ])

    # 3. Stats data
    st = store.stats()
    now_utc = datetime.now(timezone.utc)
    t_24h_ago = (now_utc - timedelta(hours=24)).strftime("%Y-%m-%dT%H:%M:%SZ")
    new_24h_count = cur.execute(
        "SELECT COUNT(*) FROM jobs WHERE first_seen >= ?", (t_24h_ago,)
    ).fetchone()[0]

    now_ist_str = now_utc.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S IST")
    stats_data = [
        {"Metric": "Total Companies", "Value": st["companies_total"]},
        {"Metric": "Disabled Companies", "Value": st["companies_disabled"]},
        {"Metric": "Errored Companies", "Value": st["companies_errored"]},
        {"Metric": "Total Jobs Tracked", "Value": st["jobs_total"]},
        {"Metric": "Open Jobs", "Value": st["jobs_open"]},
        {"Metric": "Matched Open Jobs", "Value": st["jobs_matched_open"]},
        {"Metric": "New Jobs in Last 24h", "Value": new_24h_count},
        {"Metric": "Last Export Time (IST)", "Value": now_ist_str},
    ]
    df_stats = pd.DataFrame(stats_data)

    # Write to Excel
    with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
        df_matches.to_excel(writer, sheet_name="Matches", index=False)
        df_all_open.to_excel(writer, sheet_name="All Open", index=False)
        df_stats.to_excel(writer, sheet_name="Stats", index=False)

    # Format workbook using openpyxl
    wb = openpyxl.load_workbook(temp_path)

    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
    link_font = Font(name="Calibri", size=11, color="2563EB", underline="single")

    col_widths = {
        "Role": 45,
        "Company": 25,
        "Location": 25,
        "Country": 15,
        "Min Exp": 10,
        "Max Exp": 10,
        "Exp Signal": 15,
        "Posted": 20,
        "First Seen (IST)": 20,
        "Salary": 25,
        "Source": 12,
        "Link": 60,
        "Matched": 10,
        "Reject Reason": 25,
        "Metric": 25,
        "Value": 25,
    }

    for sheet_name in ["Matches", "All Open", "Stats"]:
        ws = wb[sheet_name]
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column > 0:
            ws.auto_filter.ref = ws.dimensions

        # Format header row
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            header_val = str(cell.value or "")
            if header_val in col_widths:
                col_letter = openpyxl.utils.get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = col_widths[header_val]

        # Format hyperlinks in Link column
        if sheet_name in ["Matches", "All Open"]:
            # Find Link column index
            link_col = None
            for col_idx in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col_idx).value == "Link":
                    link_col = col_idx
                    break
            if link_col:
                for row_idx in range(2, ws.max_row + 1):
                    c = ws.cell(row=row_idx, column=link_col)
                    if c.value and str(c.value).startswith("http"):
                        c.hyperlink = str(c.value)
                        c.font = link_font

    wb.save(temp_path)
    wb.close()

    # Atomic replace
    os.replace(temp_path, out_path)
    logger.info("Successfully exported %s with 3 sheets", out_path)
